# -*- coding: utf-8 -*-
"""DEV_SUMMARY.md + PRD.md + GAME_DESIGN.md -> 하나의 .docx
사용: python build_docx.py 출력파일.docx
"""
import re, sys, os
from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

FONT = 'Malgun Gothic'
MONO = 'Consolas'
ACCENT = RGBColor(0x8B, 0x5C, 0x2E)
GREY = RGBColor(0x66, 0x66, 0x66)

def font(run, name=FONT, size=None, bold=None, italic=None, color=None):
    run.font.name = name
    rpr = run._element.get_or_add_rPr()
    rf = rpr.find(qn('w:rFonts'))
    if rf is None:
        rf = OxmlElement('w:rFonts'); rpr.append(rf)
    for k in ('w:ascii', 'w:hAnsi', 'w:eastAsia', 'w:cs'):
        rf.set(qn(k), name)
    if size: run.font.size = Pt(size)
    if bold is not None: run.font.bold = bold
    if italic is not None: run.font.italic = italic
    if color is not None: run.font.color.rgb = color

def shade(cell, hex_fill):
    tcPr = cell._element.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear'); shd.set(qn('w:color'), 'auto'); shd.set(qn('w:fill'), hex_fill)
    tcPr.append(shd)

def clean(text):
    text = text.replace('<br>', '\n')
    text = re.sub(r'<[^>]+>', '', text)
    text = re.sub(r'\[([^\]]+)\]\(([^)]+)\)', r'\1 (\2)', text)
    return text

INLINE = re.compile(r'(\*\*[^*]+\*\*|`[^`]+`)')

def add_inline(par, text, size=10.5, base_bold=False, color=None):
    text = clean(text)
    for part in INLINE.split(text):
        if not part: continue
        if part.startswith('**') and part.endswith('**'):
            r = par.add_run(part[2:-2]); font(r, size=size, bold=True, color=color)
        elif part.startswith('`') and part.endswith('`'):
            r = par.add_run(part[1:-1]); font(r, name=MONO, size=size - 0.5, color=RGBColor(0x1F, 0x3A, 0x5F))
        else:
            r = par.add_run(part); font(r, size=size, bold=base_bold or None, color=color)

def add_heading(doc, text, level):
    h = doc.add_heading('', level=level)
    r = h.add_run(clean(text))
    sizes = {1: 20, 2: 15, 3: 12.5}
    font(r, size=sizes.get(level, 12), bold=True, color=ACCENT if level == 1 else RGBColor(0x2B, 0x2B, 0x2B))
    h.paragraph_format.space_before = Pt(18 if level == 1 else 12)
    h.paragraph_format.space_after = Pt(6)
    return h

def add_table(doc, rows):
    cells = [[c.strip() for c in r.strip().strip('|').split('|')] for r in rows]
    cells = [r for r in cells if not all(re.fullmatch(r':?-{2,}:?', c or '---') for c in r)]
    if not cells: return
    ncol = max(len(r) for r in cells)
    t = doc.add_table(rows=len(cells), cols=ncol)
    t.style = 'Table Grid'
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, row in enumerate(cells):
        for j in range(ncol):
            val = row[j] if j < len(row) else ''
            cell = t.cell(i, j)
            cell.text = ''
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            lines = clean(val).split('\n')
            for li, ln in enumerate(lines):
                if li: p = cell.add_paragraph(); p.paragraph_format.space_after = Pt(0)
                add_inline(p, ln, size=9, base_bold=(i == 0))
            if i == 0: shade(cell, 'EFE6D8')
    doc.add_paragraph().paragraph_format.space_after = Pt(2)

def add_code(doc, lines):
    for ln in lines:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.left_indent = Cm(0.5)
        r = p.add_run(ln if ln else ' ')
        font(r, name=MONO, size=8.5, color=RGBColor(0x1F, 0x3A, 0x5F))
        pPr = p._p.get_or_add_pPr()
        shd = OxmlElement('w:shd'); shd.set(qn('w:val'), 'clear'); shd.set(qn('w:fill'), 'F4F1EC'); pPr.append(shd)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)

def render_md(doc, md_text, level_offset=0):
    lines = md_text.splitlines()
    i = 0
    while i < len(lines):
        ln = lines[i]
        s = ln.strip()
        if not s or s == '---':
            i += 1; continue
        if s.startswith('```'):
            j = i + 1; buf = []
            while j < len(lines) and not lines[j].strip().startswith('```'):
                buf.append(lines[j]); j += 1
            add_code(doc, buf); i = j + 1; continue
        m = re.match(r'^(#{1,4})\s+(.*)', s)
        if m:
            add_heading(doc, m.group(2), min(len(m.group(1)) + level_offset, 4)); i += 1; continue
        if s.startswith('|'):
            j = i; buf = []
            while j < len(lines) and lines[j].strip().startswith('|'):
                buf.append(lines[j]); j += 1
            add_table(doc, buf); i = j; continue
        if s.startswith('>'):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.8)
            add_inline(p, s.lstrip('> ').strip(), size=10, color=GREY)
            for r in p.runs: r.font.italic = True
            i += 1; continue
        m = re.match(r'^(\s*)[-*]\s+(.*)', ln)
        if m:
            indent = len(m.group(1).replace('\t', '  '))
            text = m.group(2)
            text = text.replace('[ ]', '☐', 1).replace('[x]', '☑', 1)
            p = doc.add_paragraph(style='List Bullet 2' if indent >= 2 else 'List Bullet')
            p.paragraph_format.space_after = Pt(2)
            add_inline(p, text); i += 1; continue
        m = re.match(r'^\s*\d+\.\s+(.*)', ln)
        if m:
            p = doc.add_paragraph(style='List Number')
            p.paragraph_format.space_after = Pt(2)
            add_inline(p, m.group(1)); i += 1; continue
        # 일반 문단 (연속 줄은 합침)
        j = i; buf = []
        while j < len(lines) and lines[j].strip() and not re.match(r'^(#{1,4}\s|\||>|```|\s*[-*]\s|\s*\d+\.\s)', lines[j]):
            buf.append(lines[j].strip()); j += 1
        p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(6)
        add_inline(p, ' '.join(buf)); i = j

def page_break(doc):
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

def cover(doc):
    for _ in range(6): doc.add_paragraph()
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run('꿈의 탈출'); font(r, size=36, bold=True, color=ACCENT)
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run('Dream Escape — 개발 정리 · PRD · 게임기획서'); font(r, size=14, color=GREY)
    doc.add_paragraph()
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run('웹 브라우저용 익스트랙션 어드벤처 · 프로토타입 v0.1.8'); font(r, size=11)
    for _ in range(8): doc.add_paragraph()
    info = [('작성일', '2026-09-25'), ('개발일', '2026-09-23'),
            ('저장소', 'https://github.com/duros221/dream-escape'),
            ('플레이', 'https://duros221.github.io/dream-escape/'),
            ('구성', 'Ⅰ. 개발 정리   Ⅱ. PRD   Ⅲ. 게임기획서')]
    t = doc.add_table(rows=len(info), cols=2); t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, (k, v) in enumerate(info):
        c0, c1 = t.cell(i, 0), t.cell(i, 1)
        c0.text = ''; c1.text = ''
        r = c0.paragraphs[0].add_run(k); font(r, size=10, bold=True, color=GREY)
        r = c1.paragraphs[0].add_run(v); font(r, size=10)
        c0.width = Cm(3); c1.width = Cm(10)

def toc(doc):
    add_heading(doc, '목차', 1)
    p = doc.add_paragraph()
    r = p.add_run()
    fld = OxmlElement('w:fldSimple')
    fld.set(qn('w:instr'), 'TOC \\o "1-2" \\h \\z \\u')
    t = OxmlElement('w:r'); tt = OxmlElement('w:t'); tt.text = '(Word에서 F9로 목차를 갱신하세요)'; t.append(tt); fld.append(t)
    r._element.addnext(fld)

def main(out):
    here = os.path.dirname(os.path.abspath(__file__))
    read = lambda n: open(os.path.join(here, n), encoding='utf-8').read()
    doc = Document()
    sec = doc.sections[0]
    sec.left_margin = sec.right_margin = Cm(2.2); sec.top_margin = sec.bottom_margin = Cm(2.0)
    st = doc.styles['Normal']; st.font.name = FONT; st.font.size = Pt(10.5)
    st.element.rPr.rFonts.set(qn('w:eastAsia'), FONT)
    cover(doc); page_break(doc)
    toc(doc); page_break(doc)
    add_heading(doc, 'Ⅰ. 개발 정리', 1)
    render_md(doc, read('DEV_SUMMARY.md').split('\n', 1)[1], level_offset=0)
    page_break(doc)
    add_heading(doc, 'Ⅱ. PRD (제품 요구사항 정의서)', 1)
    render_md(doc, read('PRD.md').split('\n', 1)[1])
    page_break(doc)
    add_heading(doc, 'Ⅲ. 게임기획서', 1)
    render_md(doc, read('GAME_DESIGN.md').split('\n', 1)[1])
    # 바닥글 페이지 번호
    for s in doc.sections:
        fp = s.footer.paragraphs[0]; fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = fp.add_run(); font(r, size=9, color=GREY)
        for tag, txt in (('begin', None), (None, 'PAGE'), ('end', None)):
            if tag:
                e = OxmlElement('w:fldChar'); e.set(qn('w:fldCharType'), tag); r._element.append(e)
            else:
                e = OxmlElement('w:instrText'); e.set(qn('xml:space'), 'preserve'); e.text = txt; r._element.append(e)
    doc.save(out)
    print('saved', out)

if __name__ == '__main__':
    main(sys.argv[1] if len(sys.argv) > 1 else os.path.join(os.path.dirname(os.path.abspath(__file__)), 'DreamEscape_개발정리.docx'))
